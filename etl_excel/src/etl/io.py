from __future__ import annotations

import os
from typing import List, Optional, Dict, Any
from openpyxl import load_workbook


SUPPORTED_FILE_SUFFIXES = {".xlsx", ".xls", ".xlsm"}


def discover_excel_files(path: str, *, recursive: bool = False) -> List[str]:
    if os.path.isfile(path):
        _, ext = os.path.splitext(path)
        return [path] if ext.lower() in SUPPORTED_FILE_SUFFIXES else []
    found: List[str] = []
    if recursive:
        for dirpath, _, filenames in os.walk(path):
            for name in filenames:
                _, ext = os.path.splitext(name)
                if ext.lower() in SUPPORTED_FILE_SUFFIXES:
                    found.append(os.path.join(dirpath, name))
    else:
        for name in os.listdir(path):
            file_path = os.path.join(path, name)
            if os.path.isfile(file_path):
                _, ext = os.path.splitext(name)
                if ext.lower() in SUPPORTED_FILE_SUFFIXES:
                    found.append(file_path)
    return sorted(found)


def read_excel_rows(path: str, sheet_name: Optional[str] = None) -> List[Dict[str, Any]]:
    wb = load_workbook(path, data_only=True, read_only=True)
    ws = wb[sheet_name] if sheet_name else wb.worksheets[0]
    rows_iter = ws.iter_rows(values_only=True)
    try:
        headers = next(rows_iter)
    except StopIteration:
        return []
    headers = [str(h).strip() if h is not None else "" for h in headers]
    records: List[Dict[str, Any]] = []
    for row in rows_iter:
        record = {headers[i]: row[i] for i in range(len(headers))}
        # skip completely empty rows
        if any(v is not None and str(v).strip() != "" for v in record.values()):
            records.append(record)
    return records

